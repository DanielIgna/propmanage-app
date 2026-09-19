"""Faza 1 — HartaBlocuri source ledger. No Building / import / parser changes."""
from __future__ import annotations

import asyncio
from pathlib import Path

import pytest
from openpyxl import Workbook

from hartablocuri_import import COL, _source_record_id, parse_workbook
from hartablocuri_source_ledger import (
    COLLECTION,
    STATUS_DE_ADAUGAT,
    STATUS_EMPTY,
    STATUS_HEADER_OR_META,
    STATUS_IMPORTED,
    STATUS_MISSING_COORDINATES,
    STATUS_MISSING_NAME,
    STATUS_MISSING_NAME_AND_COORDINATES,
    STATUS_STERGE,
    classify_source_status,
    compute_source_snapshot_id,
    ingest_source_ledger,
    iter_hartablocuri_source_rows,
    parse_eligible_for_row,
    raw_row_from_tuple,
)

REAL_XLSX = Path(__file__).resolve().parent.parent / "data" / "hartablocuri_cluj.xlsx"
NCOLS = COL["poze"] + 2


def _run(coro):
    return asyncio.run(coro)


def _blank_row():
    return [None] * NCOLS


def _set(row, key, value):
    row[COL[key]] = value
    return row


def _row(**fields):
    row = _blank_row()
    for k, v in fields.items():
        _set(row, k, v)
    return tuple(row)


def _write_xlsx(path: Path, rows: list[tuple]):
    wb = Workbook()
    ws = wb.active
    ws.title = "Detalii blocuri"
    for r_i, row in enumerate(rows, start=1):
        for c_i, val in enumerate(row, start=1):
            if val is not None:
                ws.cell(r_i, c_i, val)
    # Force sheet width past COL["poze"] without filling empty rows.
    ws.cell(2, NCOLS, "Poze")
    wb.save(path)
    wb.close()


class _WriteGuard:
    def __init__(self, name):
        self.name = name
        self.calls = []

    def _boom(self, *args, **kwargs):
        self.calls.append((args, kwargs))
        raise AssertionError(f"Faza 1 must not mutate {self.name}")

    insert_one = insert_many = update_one = update_many = replace_one = _boom
    delete_one = delete_many = bulk_write = find_one_and_update = _boom
    find_one_and_replace = find_one_and_delete = _boom
    create_index = _boom


class _FakeColl:
    def __init__(self):
        self.docs = {}
        self.indexes = []

    async def create_index(self, keys, **kwargs):
        self.indexes.append((keys, kwargs))
        return kwargs.get("name") or "idx"

    async def bulk_write(self, ops, ordered=False):
        upserted = modified = matched = 0
        for op in ops:
            filt = getattr(op, "filter", None) or op._filter
            doc = getattr(op, "replacement", None)
            if doc is None:
                doc = getattr(op, "_doc", None)
            key = (filt["source_snapshot_id"], filt["excel_row"])
            if key in self.docs:
                matched += 1
                if self.docs[key] != doc:
                    modified += 1
            else:
                upserted += 1
            self.docs[key] = doc

        class R:
            pass
        r = R()
        r.upserted_count = upserted
        r.modified_count = modified
        r.matched_count = matched
        return r

    async def count_documents(self, q):
        snap = q.get("source_snapshot_id")
        n = 0
        for (s, _row), _doc in self.docs.items():
            if snap is None or s == snap:
                n += 1
        return n


class _FakeDB:
    def __init__(self):
        self.hartablocuri_source_records = _FakeColl()
        self.buildings = _WriteGuard("buildings")
        self.properties = _WriteGuard("properties")
        self.import_batches = _WriteGuard("import_batches")

    def __getitem__(self, name):
        if name == COLLECTION:
            return self.hartablocuri_source_records
        raise KeyError(name)


def _fixture_rows():
    empty = tuple(_blank_row())
    header = _row(nume="Nume", lat="Latitudine", lng="Longitudine", adresa="Adresă", era="Eră")
    counta = _row(nume=3881, lat=3752, lng=3752, adresa=3817, era=3884)
    meta = _row(nume="Compusă de Teoalida © www.teoalida.ro")
    imported = _row(
        nume="Bloc A1", lat=46.76485, lng=23.603, adresa="Aleea Muscel nr. 19",
        oras="Cluj-Napoca", uat="municipiu Cluj-Napoca > cartier Gheorgheni",
        era="comunist 1950-1969", scari=1, apartamente=10, an_finalizare="1965",
        structura="cărămidă",
    )
    sterge = _row(
        nume="C4", lat=46.7465, lng=23.5292, adresa="Strada Valea Gârboului nr. 156",
        era="STERGE", scari=1, apartamente=8, an_finalizare="2023 estimare Andrei",
    )
    de_ad = _row(
        nume="Bloc L3", lat=46.7765, lng=23.6375, adresa="Strada Tulcea nr. 26",
        era="DE ADĂUGAT", an_finalizare="1987 estimare Andrei",
        structura="panouri prefabricate",
        planuri='<a href="//www.hartablocuri.ro/planuri/x.png">p</a>',
    )
    numeric = _row(
        nume=15, lat=46.7674, lng=23.6197, adresa="Strada Liviu Rebreanu nr. 15",
        era="comunist 1950-1969", scari=1, apartamente=44,
        planuri='<a href="//www.hartablocuri.ro/planuri/rebreanu.png">p</a>',
    )
    missing_xy = _row(nume="Bloc fără coordonate", adresa="Strada Exemplu nr. 1", era="comunist 1950-1969")
    missing_both = _row(era="DE ADĂUGAT", oras="Cluj-Napoca")
    return [
        empty, header, counta, meta, imported, sterge, de_ad, numeric, missing_xy, missing_both,
    ]


# ── classification ──────────────────────────────────────────────


def test_empty_classified():
    assert classify_source_status(tuple(_blank_row())) == STATUS_EMPTY


def test_header_and_meta_classified():
    assert classify_source_status(_row(nume="Nume", lat="Latitudine")) == STATUS_HEADER_OR_META
    assert classify_source_status(_row(nume="Name", lat="Latitude")) == STATUS_HEADER_OR_META
    assert classify_source_status(_row(nume=3881, lat=3752, lng=3752)) == STATUS_HEADER_OR_META
    assert classify_source_status(_row(nume="Compusă de Teoalida © www.teoalida.ro")) == STATUS_HEADER_OR_META


def test_sterge_kept_as_source_status_not_delete():
    st = classify_source_status(_row(
        nume="C4", lat=46.7465, lng=23.5292, era="STERGE", adresa="Strada X nr. 1",
    ))
    assert st == STATUS_STERGE
    row = _row(nume="C4", lat=46.7465, lng=23.5292, era="STERGE", adresa="Strada X nr. 1")
    assert parse_eligible_for_row(row) is False


def test_de_adaugat_kept_as_source_status_not_delete():
    st = classify_source_status(_row(
        nume="Bloc L3", lat=46.7765, lng=23.6375, era="DE ADĂUGAT", adresa="Strada Tulcea nr. 26",
    ))
    assert st == STATUS_DE_ADAUGAT
    assert parse_eligible_for_row(_row(
        nume="Bloc L3", lat=46.7765, lng=23.6375, era="DE ADĂUGAT", adresa="Strada Tulcea nr. 26",
    )) is False


def test_missing_coordinates_and_name():
    assert classify_source_status(_row(nume="Bloc X", adresa="Strada Y nr. 1")) == STATUS_MISSING_COORDINATES
    assert classify_source_status(_row(era="DE ADĂUGAT")) == STATUS_MISSING_NAME_AND_COORDINATES
    assert classify_source_status(_row(
        nume=15, lat=46.7674, lng=23.6197, adresa="Strada Liviu Rebreanu nr. 15",
        era="comunist 1950-1969",
    )) == STATUS_MISSING_NAME


def test_numeric_name_raw_and_nume_as_text(tmp_path):
    rows = _fixture_rows()
    path = tmp_path / "hartablocuri_cluj.xlsx"
    _write_xlsx(path, rows)
    docs = list(iter_hartablocuri_source_rows(str(path)))
    numeric = docs[7]
    assert numeric["source_status"] == STATUS_MISSING_NAME
    assert numeric["raw_row"]["2"] == 15
    assert numeric["parsed"]["nume"] is None
    assert numeric["parsed"]["nume_as_text"] == "15"
    assert numeric["promoted_to_building"] is False
    assert numeric["valuable_excluded"] is True


def test_valuable_excluded_not_hardcoded(tmp_path):
    rows = _fixture_rows()
    path = tmp_path / "hartablocuri_cluj.xlsx"
    _write_xlsx(path, rows)
    docs = list(iter_hartablocuri_source_rows(str(path)))
    by_status = {d["source_status"]: d for d in docs}
    assert by_status[STATUS_IMPORTED]["valuable_excluded"] is False
    assert by_status[STATUS_STERGE]["valuable_excluded"] is True
    assert by_status[STATUS_DE_ADAUGAT]["valuable_excluded"] is True
    assert by_status[STATUS_EMPTY]["valuable_excluded"] is False
    assert by_status[STATUS_HEADER_OR_META]["valuable_excluded"] is False
    assert by_status[STATUS_MISSING_COORDINATES]["valuable_excluded"] is False


def test_promoted_flags_and_no_building_pointer(tmp_path):
    path = tmp_path / "hartablocuri_cluj.xlsx"
    _write_xlsx(path, _fixture_rows())
    docs = list(iter_hartablocuri_source_rows(str(path)))
    imported = [d for d in docs if d["source_status"] == STATUS_IMPORTED]
    others = [d for d in docs if d["source_status"] != STATUS_IMPORTED]
    assert len(imported) == 1
    assert imported[0]["parse_eligible"] is True
    assert imported[0]["promoted_to_building"] is True
    assert imported[0]["building_id"] is None
    assert imported[0]["link_status"] == "none"
    assert all(d["promoted_to_building"] is False for d in others)
    assert all(d["building_id"] is None for d in docs)
    sterge = next(d for d in docs if d["source_status"] == STATUS_STERGE)
    de_ad = next(d for d in docs if d["source_status"] == STATUS_DE_ADAUGAT)
    assert sterge["promoted_to_building"] is False
    assert de_ad["promoted_to_building"] is False


def test_existing_source_record_id_matches_parse_workbook(tmp_path):
    path = tmp_path / "hartablocuri_cluj.xlsx"
    _write_xlsx(path, _fixture_rows())
    parsed = parse_workbook(str(path))
    docs = [d for d in iter_hartablocuri_source_rows(str(path)) if d["parse_eligible"]]
    assert len(parsed) == 1
    assert len(docs) == 1
    assert docs[0]["source_record_id"] == parsed[0]["source_record_id"]
    rec = {
        "nume": parsed[0]["nume"], "adresa": parsed[0]["adresa"],
        "lat": parsed[0]["lat"], "lng": parsed[0]["lng"],
    }
    assert docs[0]["source_record_id"] == _source_record_id(rec)


def test_ingest_rerun_zero_duplicates_and_no_building_writes(tmp_path):
    path = tmp_path / "hartablocuri_cluj.xlsx"
    _write_xlsx(path, _fixture_rows())
    db = _FakeDB()
    s1 = _run(ingest_source_ledger(db, str(path)))
    n1 = len(db.hartablocuri_source_records.docs)
    s2 = _run(ingest_source_ledger(db, str(path)))
    n2 = len(db.hartablocuri_source_records.docs)
    assert n1 == n2 == s1["physical_rows"] == s2["stored"]
    assert s2["upserted"] == 0
    assert db.buildings.calls == []
    assert db.properties.calls == []
    assert db.import_batches.calls == []
    idx_names = [kwargs.get("name") for _keys, kwargs in db.hartablocuri_source_records.indexes]
    assert "snapshot_excel_row_unique" in idx_names
    assert "building_id_sparse" in idx_names


def test_raw_row_keeps_unmapped_columns():
    row = list(_blank_row())
    row[21] = "coloană nemapată"
    raw = raw_row_from_tuple(tuple(row))
    assert raw["22"] == "coloană nemapată"
    assert "2" in raw


@pytest.mark.skipif(not REAL_XLSX.exists(), reason="hartablocuri_cluj.xlsx missing")
def test_real_workbook_ids_match_parse_workbook():
    parsed = parse_workbook(str(REAL_XLSX))
    imported = [d for d in iter_hartablocuri_source_rows(str(REAL_XLSX)) if d["parse_eligible"]]
    assert len(parsed) == 3406
    assert len(imported) == 3406
    assert {d["source_record_id"] for d in imported} == {r["source_record_id"] for r in parsed}


@pytest.mark.skipif(not REAL_XLSX.exists(), reason="hartablocuri_cluj.xlsx missing")
def test_real_workbook_audit_counts():
    docs = list(iter_hartablocuri_source_rows(str(REAL_XLSX)))
    from collections import Counter
    c = Counter(d["source_status"] for d in docs)
    assert len(docs) == 3924
    assert c[STATUS_IMPORTED] == 3406
    assert c[STATUS_STERGE] == 200
    assert c[STATUS_DE_ADAUGAT] == 63
    assert c[STATUS_EMPTY] == 6
    assert sum(d["promoted_to_building"] for d in docs) == 3406
    valuable = sum(1 for d in docs if d["valuable_excluded"])
    assert valuable == 55
    assert all(d["building_id"] is None and d["link_status"] == "none" for d in docs)
    assert compute_source_snapshot_id(str(REAL_XLSX)).startswith("hartablocuri_cluj_detalii-blocuri_")
    # header_or_meta is extra vs the prior audit (those rows were missing_*).
    assert c[STATUS_HEADER_OR_META] >= 1
    assert c[STATUS_MISSING_NAME] + c[STATUS_HEADER_OR_META] >= 85
    assert parse_workbook(str(REAL_XLSX))  # parser still 3406 via previous assert
